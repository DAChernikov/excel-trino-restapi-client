from __future__ import annotations

import argparse
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

from trino_excel_client.template import REQUIRED_TABLE_NAMES, client_sheet_names


def _run(command: list[str]) -> None:
    print(">", " ".join(command))
    subprocess.run(command, check=True)


def _table_names(workbook_path: Path) -> set[str]:
    workbook = load_workbook(workbook_path, read_only=False, keep_vba=workbook_path.suffix.lower() == ".xlsm")
    names: set[str] = set()
    for worksheet in workbook.worksheets:
        names.update(worksheet.tables.keys())
    return names


def _assert_ui_contract(workbook_path: Path, sheet_prefix: str) -> None:
    workbook = load_workbook(workbook_path, read_only=False, keep_vba=workbook_path.suffix.lower() == ".xlsm")
    names = client_sheet_names(sheet_prefix)

    missing_sheets = [sheet_name for sheet_name in names.all if sheet_name not in workbook.sheetnames]
    if missing_sheets:
        raise AssertionError(f"Missing sheets: {missing_sheets}")

    missing_tables = sorted(set(REQUIRED_TABLE_NAMES) - _table_names(workbook_path))
    if missing_tables:
        raise AssertionError(f"Missing tables: {missing_tables}")

    if workbook[names.query]["A5"].value != "select * from your_catalog.your_schema.your_table limit 100":
        raise AssertionError("Unexpected SQL template value in Trino Query!A5")

    if workbook[names.result]["A1"].value != "Trino REST API Client - Result":
        raise AssertionError("Trino Result sheet is not present or malformed")

    if workbook[names.config]["B18"].value != "TRUE":
        raise AssertionError("Default SQL limit switch must be TRUE")

    if str(workbook[names.config]["B19"].value) != "100000":
        raise AssertionError("max_result_rows default must be 100000")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Windows smoke test for the Trino Excel client. "
            "Requires desktop Excel and pywin32 for --backend com."
        )
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("build") / "smoke",
        help="Directory for generated smoke-test workbooks.",
    )
    parser.add_argument(
        "--sheet-prefix",
        default="Trino",
        help='Sheet prefix to use. Default: "Trino".',
    )
    parser.add_argument(
        "--visible",
        action="store_true",
        help="Show Excel while COM backend runs.",
    )
    parser.add_argument(
        "--skip-com",
        action="store_true",
        help="Only verify openpyxl UI generation. Useful on non-Windows CI.",
    )
    return parser


def main() -> None:
    args = build_parser().parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)

    openpyxl_output = args.output_dir / "Trino_UI_Preview.xlsx"
    com_output = args.output_dir / "Trino_REST_Client_COM.xlsx"

    _run(
        [
            sys.executable,
            "-m",
            "trino_excel_client.cli",
            "create",
            "--backend",
            "openpyxl",
            "--output",
            str(openpyxl_output),
            "--overwrite",
            "--sheet-prefix",
            args.sheet_prefix,
        ]
    )
    _assert_ui_contract(openpyxl_output, args.sheet_prefix)
    print(f"OK: openpyxl UI workbook contract verified: {openpyxl_output}")

    if args.skip_com:
        print("Skipped COM backend smoke test.")
        return

    command = [
        sys.executable,
        "-m",
        "trino_excel_client.cli",
        "create",
        "--backend",
        "com",
        "--output",
        str(com_output),
        "--overwrite",
        "--sheet-prefix",
        args.sheet_prefix,
    ]
    if args.visible:
        command.append("--visible")

    _run(command)
    _assert_ui_contract(com_output, args.sheet_prefix)
    print(f"OK: COM workbook UI contract verified: {com_output}")
    print("Manual release check: open the COM workbook, fill Config, refresh several times, and confirm Trino Query remains unchanged.")


if __name__ == "__main__":
    main()
