from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .template import (
    CONFIG_TABLE,
    EXTRA_CREDENTIALS_TABLE,
    HELP_TABLE,
    SQL_TABLE,
    TableSpec,
    client_sheet_names,
)

COLORS = {
    "navy": "1F3349",
    "blue": "2E75B6",
    "light_blue": "DDEBF7",
    "green": "70AD47",
    "light_green": "E2EFDA",
    "light_orange": "FCE4D6",
    "gray": "F2F2F2",
    "dark_gray": "595959",
    "white": "FFFFFF",
}


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=COLORS[color])


def _write_rows(ws, start_row: int, start_col: int, rows: Iterable[Iterable[object]]) -> None:
    for row_offset, row in enumerate(rows):
        for col_offset, value in enumerate(row):
            ws.cell(start_row + row_offset, start_col + col_offset, value=value)


def _table_ref(start_row: int, start_col: int, row_count: int, col_count: int) -> str:
    top_left = ws_cell(start_row, start_col)
    bottom_right = ws_cell(start_row + row_count - 1, start_col + col_count - 1)
    return f"{top_left}:{bottom_right}"


def ws_cell(row: int, col: int) -> str:
    letters = ""
    current = col
    while current:
        current, remainder = divmod(current - 1, 26)
        letters = chr(65 + remainder) + letters
    return f"{letters}{row}"


def _add_table(ws, start_cell: str, spec: TableSpec) -> Table:
    start_col = ws[start_cell].column
    start_row = ws[start_cell].row
    data = (spec.headers, *spec.rows)
    _write_rows(ws, start_row, start_col, data)

    ref = _table_ref(start_row, start_col, len(data), len(spec.headers))
    table = Table(displayName=spec.name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name=spec.style,
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    for row in ws[ref]:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name="Aptos", size=10)
    for cell in ws[start_row]:
        cell.font = Font(name="Aptos", size=10, bold=True, color=COLORS["white"])
        cell.fill = _fill("blue")
    return table


def _title(ws, address: str, text: str, width: int) -> None:
    cell = ws[address]
    cell.value = text
    end_col = cell.column + width - 1
    ws.merge_cells(start_row=cell.row, start_column=cell.column, end_row=cell.row, end_column=end_col)
    cell.fill = _fill("navy")
    cell.font = Font(name="Aptos", size=16, bold=True, color=COLORS["white"])
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[cell.row].height = 32


def _section(ws, address: str, text: str, width: int) -> None:
    cell = ws[address]
    cell.value = text
    end_col = cell.column + width - 1
    ws.merge_cells(start_row=cell.row, start_column=cell.column, end_row=cell.row, end_column=end_col)
    cell.fill = _fill("light_blue")
    cell.font = Font(name="Aptos", size=10, bold=True, color=COLORS["navy"])
    ws.row_dimensions[cell.row].height = 24


def _note(ws, address: str, text: str, width: int, fill: str) -> None:
    cell = ws[address]
    cell.value = text
    end_col = cell.column + width - 1
    ws.merge_cells(start_row=cell.row, start_column=cell.column, end_row=cell.row, end_column=end_col)
    cell.fill = _fill(fill)
    cell.font = Font(name="Aptos", size=10, italic=True, color=COLORS["dark_gray"])
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[cell.row].height = 46


def _delete_table_if_exists(wb, table_name: str) -> None:
    for ws in wb.worksheets:
        if table_name in ws.tables:
            del ws.tables[table_name]


def _remove_client_sheets(wb, sheet_prefix: str) -> None:
    names = client_sheet_names(sheet_prefix)
    for name in names.all:
        if name in wb.sheetnames:
            del wb[name]


def _add_config_sheet(wb, sheet_prefix: str):
    ws = wb.create_sheet(client_sheet_names(sheet_prefix).config)
    ws.sheet_properties.tabColor = COLORS["blue"]
    _title(ws, "A1", "Trino REST API Client - Config", 4)
    _note(
        ws,
        "A2",
        "Заполните колонку 'значение'. Парольные ячейки только визуально скрыты, однако не защищают ваши учетные данные, будьте осторожны.",
        4,
        "light_orange",
    )
    _section(ws, "A4", "Параметры подключения", 4)
    _add_table(ws, "A5", CONFIG_TABLE)
    ws["B8"].number_format = ";;;"

    _section(ws, "A23", "Extra credentials для источников данных", 5)
    _note(
        ws,
        "A24",
        "Используйте эту таблицу для заполнения extraCredentials параметров Trino, например gp-user/gp-password. Активные строки должны быть = TRUE для использования.",
        5,
        "light_green",
    )
    _add_table(ws, "A26", EXTRA_CREDENTIALS_TABLE)
    for row in range(27, 30):
        ws[f"E{row}"].number_format = ";;;"

    validation = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
    ws.add_data_validation(validation)
    validation.add("A27:A29")

    overflow_validation = DataValidation(type="list", formula1='"error,truncate"', allow_blank=False)
    ws.add_data_validation(overflow_validation)
    overflow_validation.add("B20")

    apply_limit_validation = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=False)
    ws.add_data_validation(apply_limit_validation)
    apply_limit_validation.add("B18")

    widths = {"A": 28, "B": 34, "C": 13, "D": 62, "E": 24, "F": 32}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A5"
    return ws


def _add_query_sheet(wb, sheet_prefix: str):
    ws = wb.create_sheet(client_sheet_names(sheet_prefix).query)
    ws.sheet_properties.tabColor = COLORS["blue"]
    _title(ws, "A1", "Trino REST API Client - Query", 4)
    _note(
        ws,
        "A2",
        "Введите SQL в одну широкую строку ниже. Результат загружается отдельно на лист Trino Result, чтобы обновление не затрагивало ввод.",
        4,
        "light_blue",
    )
    _add_table(ws, "A4", SQL_TABLE)
    ws["A5"].font = Font(name="Consolas", size=10)
    ws["A5"].alignment = Alignment(vertical="center", wrap_text=False)
    ws.row_dimensions[4].height = 22
    ws.row_dimensions[5].height = 24

    ws.column_dimensions["A"].width = 150
    ws.freeze_panes = "A5"
    return ws


def _add_result_sheet(wb, sheet_prefix: str):
    ws = wb.create_sheet(client_sheet_names(sheet_prefix).result)
    ws.sheet_properties.tabColor = COLORS["green"]
    _title(ws, "A1", "Trino REST API Client - Result", 4)
    _note(
        ws,
        "A2",
        "Сюда загружается результат Power Query TrinoResult. Не размещайте пользовательские данные ниже этой строки: область результата постоянно перезатирается.",
        4,
        "light_green",
    )
    _section(ws, "A4", "Результат", 4)
    ws["A5"] = "В preview-книге на Unix есть только UI-таблицы. Для встраивания Power Query используйте WindowsOS."
    ws["A5"].fill = _fill("gray")
    ws["A5"].font = Font(name="Aptos", size=10, italic=True, color=COLORS["dark_gray"])
    ws["A5"].alignment = Alignment(vertical="top", wrap_text=True)
    ws.column_dimensions["A"].width = 60
    ws.freeze_panes = "A5"
    return ws


def _add_help_sheet(wb, sheet_prefix: str):
    ws = wb.create_sheet(client_sheet_names(sheet_prefix).help)
    ws.sheet_properties.tabColor = COLORS["blue"]
    _title(ws, "A1", "Trino REST API Client - Help", 3)
    _add_table(ws, "A4", HELP_TABLE)
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 96
    return ws


def install_client_ui(wb, sheet_prefix: str = "Trino") -> None:
    for table_name in (
        CONFIG_TABLE.name,
        EXTRA_CREDENTIALS_TABLE.name,
        SQL_TABLE.name,
        HELP_TABLE.name,
    ):
        _delete_table_if_exists(wb, table_name)
    _remove_client_sheets(wb, sheet_prefix)
    _add_config_sheet(wb, sheet_prefix)
    _add_query_sheet(wb, sheet_prefix)
    _add_result_sheet(wb, sheet_prefix)
    _add_help_sheet(wb, sheet_prefix)


def create_workbook_ui(output_path: Path, overwrite: bool = False, sheet_prefix: str = "Trino") -> Path:
    output_path = output_path.resolve()
    if output_path.exists() and not overwrite:
        raise FileExistsError(f"File already exists: {output_path}")
    if output_path.suffix.lower() == ".xlsm":
        raise ValueError("openpyxl backend can create .xlsx only. Use Windows COM backend for new .xlsm files.")
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    del wb[wb.sheetnames[0]]
    install_client_ui(wb, sheet_prefix=sheet_prefix)
    wb.save(output_path)
    return output_path


def install_client_ui_file(
    workbook_path: Path,
    output_path: Path | None = None,
    overwrite: bool = False,
    sheet_prefix: str = "Trino",
) -> Path:
    workbook_path = workbook_path.resolve()
    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook does not exist: {workbook_path}")
    final_path = output_path.resolve() if output_path else workbook_path
    if output_path is not None and final_path.exists() and final_path != workbook_path and not overwrite:
        raise FileExistsError(f"File already exists: {final_path}")

    keep_vba = workbook_path.suffix.lower() == ".xlsm"
    wb = load_workbook(workbook_path, keep_vba=keep_vba)
    install_client_ui(wb, sheet_prefix=sheet_prefix)
    final_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(final_path)
    return final_path
