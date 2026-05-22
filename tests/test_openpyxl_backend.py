from pathlib import Path

from openpyxl import Workbook, load_workbook

from trino_excel_client.openpyxl_backend import create_workbook_ui, install_client_ui_file
from trino_excel_client.template import REQUIRED_TABLE_NAMES, client_sheet_names


def _table_names(path: Path) -> set[str]:
    workbook = load_workbook(path)
    names: set[str] = set()
    for worksheet in workbook.worksheets:
        names.update(worksheet.tables.keys())
    return names


def test_create_workbook_ui_creates_client_sheets_and_tables(tmp_path: Path) -> None:
    output = tmp_path / "client.xlsx"

    create_workbook_ui(output, overwrite=False)

    workbook = load_workbook(output)
    names = client_sheet_names()
    assert workbook.sheetnames == list(names.all)
    assert set(REQUIRED_TABLE_NAMES).issubset(_table_names(output))
    assert workbook[names.config]["A1"].value == "Trino REST API Client - Config"
    assert workbook[names.config]["A5"].value == "Параметр"
    assert "однако не защищают" in workbook[names.config]["A2"].value
    assert "extraCredentials параметров Trino" in workbook[names.config]["A24"].value
    assert workbook[names.config]["B7"].number_format != ";;;"
    assert workbook[names.config]["B8"].number_format == ";;;"
    assert workbook[names.config]["B11"].value == "Microsoft Excel PowerQuery client"
    assert workbook[names.config]["A17"].value == "default_query_limit_rows"
    assert workbook[names.config]["B17"].value == "100000"
    assert workbook[names.config]["A18"].value == "apply_default_query_limit"
    assert workbook[names.config]["B18"].value == "TRUE"
    assert workbook[names.config]["A19"].value == "max_result_rows"
    assert workbook[names.config]["B19"].value == "100000"
    assert workbook[names.config]["A20"].value == "result_overflow_behavior"
    assert workbook[names.config]["B20"].value == "error"
    assert workbook[names.config]["D6"].value.startswith("URL координатора")
    assert workbook[names.query]["A4"].value == "SQL запрос"
    assert workbook[names.query]["A5"].value == "select * from your_catalog.your_schema.your_table limit 100"
    assert workbook[names.query]["A5"].alignment.wrap_text is not True
    assert workbook[names.result]["A1"].value == "Trino REST API Client - Result"
    assert "область результата постоянно перезатирается" in workbook[names.result]["A2"].value
    assert workbook[names.result]["A5"].value == "В preview-книге на Unix есть только UI-таблицы. Для встраивания Power Query используйте WindowsOS."


def test_install_client_ui_preserves_existing_sheet(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    output = tmp_path / "installed.xlsx"

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Existing"
    worksheet["A1"] = "keep me"
    workbook.save(source)

    install_client_ui_file(source, output_path=output, overwrite=False)

    installed = load_workbook(output)
    assert "Existing" in installed.sheetnames
    assert installed["Existing"]["A1"].value == "keep me"
    assert set(client_sheet_names().all).issubset(installed.sheetnames)
    assert set(REQUIRED_TABLE_NAMES).issubset(_table_names(output))
