from __future__ import annotations

import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook

from trino_excel_client import excel_com
from trino_excel_client.m_code import M_QUERIES
from trino_excel_client.template import REQUIRED_TABLE_NAMES, client_sheet_names


class FakePythonCom:
    def __init__(self) -> None:
        self.initialized = 0
        self.uninitialized = 0

    def CoInitialize(self) -> None:
        self.initialized += 1

    def CoUninitialize(self) -> None:
        self.uninitialized += 1


class FakeRange:
    def __init__(self, address: str, worksheet=None) -> None:
        self.address = address
        self.Worksheet = worksheet
        self.Value = None
        self._formula = None
        self.FormulaLocal = None
        self.WrapText = None
        self.cleared = False
        self.Left = 10
        self.Top = 10
        self.Interior = type("Interior", (), {"Color": None})()
        self.FormatConditions = FakeFormatConditions()

    @property
    def Formula(self):
        return self._formula

    @Formula.setter
    def Formula(self, value) -> None:
        self._formula = value
        self.FormulaLocal = str(value).replace("INDEX", "LOCAL_INDEX").replace("MATCH", "LOCAL_MATCH")

    def Clear(self) -> None:
        self.cleared = True
        self._formula = None
        self.FormulaLocal = None


class FakeFormatCondition:
    def __init__(self, formula: str) -> None:
        self.Formula1 = formula
        self.NumberFormat = None
        self.StopIfTrue = None


class FakeFormatConditions:
    def __init__(self) -> None:
        self.items: list[FakeFormatCondition] = []
        self.deleted = False

    def Delete(self) -> None:
        self.deleted = True
        self.items.clear()

    def Add(self, *, Type, Formula1: str) -> FakeFormatCondition:
        condition = FakeFormatCondition(Formula1)
        self.items.append(condition)
        return condition


class FakeQueryTable:
    def __init__(self, workbook_connection=None) -> None:
        self.CommandType = None
        self.CommandText = None
        self.BackgroundQuery = None
        self.RefreshOnFileOpen = None
        self.RefreshPeriod = None
        self.EnableRefresh = None
        self.SaveData = None
        self.PreserveFormatting = None
        self.PreserveColumnInfo = None
        self.AdjustColumnWidth = None
        self.RefreshStyle = None
        self.WorkbookConnection = workbook_connection


class FakeListObject:
    def __init__(self, workbook_connection=None) -> None:
        self.Name = None
        self.TableStyle = None
        self.QueryTable = FakeQueryTable(workbook_connection)
        self.deleted = False

    def Delete(self) -> None:
        self.deleted = True


class FakeListObjects:
    def __init__(self, workbook_connection=None) -> None:
        self.add_calls: list[dict[str, object]] = []
        self.fail_add = False
        self.result = FakeListObject(workbook_connection)

    def __call__(self, name: str) -> FakeListObject:
        raise KeyError(name)

    def Add(self, source_type, connection_string, link_source, has_headers, destination):
        if self.fail_add:
            raise RuntimeError("cannot create result table")
        self.add_calls.append(
            {
                "source_type": source_type,
                "connection_string": connection_string,
                "link_source": link_source,
                "has_headers": has_headers,
                "destination": destination,
            }
        )
        return self.result


class FakeWorksheet:
    def __init__(self, name: str, workbook_connection=None) -> None:
        self.name = name
        self.ListObjects = FakeListObjects(workbook_connection)
        self.Shapes = FakeShapes()
        self.ranges: dict[str, FakeRange] = {}
        self.Visible = None

    def Range(self, address: str) -> FakeRange:
        if address not in self.ranges:
            self.ranges[address] = FakeRange(address, self)
        return self.ranges[address]


class FakeWorksheets:
    def __init__(self, sheet_prefix: str = "Trino", result_connection=None, schema_connection=None) -> None:
        self.query_sheet = FakeWorksheet(client_sheet_names(sheet_prefix).query)
        self.result_sheet = FakeWorksheet(client_sheet_names(sheet_prefix).result, result_connection)
        self.schema_sheet = FakeWorksheet(client_sheet_names(sheet_prefix).schema, schema_connection)
        self.by_name = {
            self.query_sheet.name: self.query_sheet,
            self.result_sheet.name: self.result_sheet,
            self.schema_sheet.name: self.schema_sheet,
        }

    def __call__(self, name: str) -> FakeWorksheet:
        return self.by_name[name]


class FakeQueries:
    def __init__(self) -> None:
        self.added: list[dict[str, str]] = []

    def __call__(self, name: str):
        raise KeyError(name)

    def Add(self, *, Name: str, Formula: str, Description: str) -> None:
        self.added.append(
            {
                "Name": Name,
                "Formula": Formula,
                "Description": Description,
            }
        )


class FakeName:
    def __init__(self, name: str) -> None:
        self.name = name
        self.deleted = False

    def Delete(self) -> None:
        self.deleted = True


class FakeNames:
    def __init__(self) -> None:
        self.added: list[dict[str, str]] = []
        self._names: dict[str, FakeName] = {}

    def __call__(self, name: str) -> FakeName:
        return self._names[name]

    def Add(self, *, Name: str, RefersTo: str) -> None:
        self.added.append({"Name": Name, "RefersTo": RefersTo})
        self._names[Name] = FakeName(Name)


class FakeCodeModule:
    def __init__(self) -> None:
        self.code = ""

    def AddFromString(self, code: str) -> None:
        self.code = code


class FakeVBComponent:
    def __init__(self) -> None:
        self.Name = ""
        self.CodeModule = FakeCodeModule()


class FakeVBComponents:
    def __init__(self) -> None:
        self.components: dict[str, FakeVBComponent] = {}
        self.added: list[FakeVBComponent] = []
        self.removed: list[FakeVBComponent] = []

    def __call__(self, name: str) -> FakeVBComponent:
        return self.components[name]

    def Add(self, component_type: int) -> FakeVBComponent:
        assert component_type == 1
        component = FakeVBComponent()
        self.added.append(component)
        return component

    def Remove(self, component: FakeVBComponent) -> None:
        self.removed.append(component)


class FakeVBProject:
    def __init__(self) -> None:
        self.VBComponents = FakeVBComponents()


class FakeShapeTextCharacters:
    def __init__(self) -> None:
        self.Text = ""


class FakeShapeTextFrame:
    def __init__(self) -> None:
        self.characters = FakeShapeTextCharacters()

    def Characters(self) -> FakeShapeTextCharacters:
        return self.characters


class FakeShape:
    def __init__(self) -> None:
        self.Name = None
        self.OnAction = None
        self.TextFrame = FakeShapeTextFrame()
        self.deleted = False

    def Delete(self) -> None:
        self.deleted = True


class FakeShapes:
    def __init__(self) -> None:
        self.items: list[FakeShape] = []
        self.Count = 0

    def __call__(self, index: int) -> FakeShape:
        return self.items[index - 1]

    def AddShape(self, shape_type, left, top, width, height) -> FakeShape:
        shape = FakeShape()
        self.items.append(shape)
        self.Count = len(self.items)
        return shape


class FakeConnectionChild:
    def __init__(self) -> None:
        self.BackgroundQuery = None
        self.RefreshOnFileOpen = None
        self.RefreshPeriod = None
        self.EnableRefresh = None


class FakeConnection:
    def __init__(self, name: str) -> None:
        self.Name = name
        self.RefreshWithRefreshAll = None
        self.RefreshOnFileOpen = None
        self.OLEDBConnection = FakeConnectionChild()


class FakeConnections:
    def __init__(self, connections: list[FakeConnection]) -> None:
        self._connections = connections
        self.Count = len(connections)

    def __call__(self, index: int) -> FakeConnection:
        return self._connections[index - 1]


class FakeWorkbook:
    def __init__(self, path: Path) -> None:
        self.path = path
        self.Name = path.name
        self.Queries = FakeQueries()
        self.Names = FakeNames()
        self.VBProject = FakeVBProject()
        self.helper_connection = FakeConnection("Query - qConfig")
        self.result_connection = FakeConnection("Query - TrinoResult")
        self.schema_connection = FakeConnection("Query - TrinoResultSchema")
        self.Connections = FakeConnections([self.helper_connection, self.schema_connection, self.result_connection])
        self.Worksheets = FakeWorksheets(
            result_connection=self.result_connection,
            schema_connection=self.schema_connection,
        )
        self.saved = False
        self.save_as_calls: list[dict[str, object]] = []
        self.close_calls: list[bool] = []

    def Save(self) -> None:
        self.saved = True

    def SaveAs(self, path: str, FileFormat: int) -> None:
        self.save_as_calls.append({"path": path, "FileFormat": FileFormat})
        shutil.copyfile(self.path, path)
        self.path = Path(path)
        self.Name = self.path.name
        self.saved = True

    def Close(self, SaveChanges: bool) -> None:
        self.close_calls.append(SaveChanges)


class FakeWorkbooks:
    def __init__(self) -> None:
        self.opened_paths: list[str] = []
        self.last_workbook: FakeWorkbook | None = None

    def Open(self, path: str) -> FakeWorkbook:
        self.opened_paths.append(path)
        self.last_workbook = FakeWorkbook(Path(path))
        return self.last_workbook


class FakeExcel:
    def __init__(self) -> None:
        self.Visible = None
        self.DisplayAlerts = None
        self.Workbooks = FakeWorkbooks()
        self.quit_called = False

    def Quit(self) -> None:
        self.quit_called = True


class FakeWin32:
    def __init__(self) -> None:
        self.excel = FakeExcel()

    def DispatchEx(self, name: str) -> FakeExcel:
        assert name == "Excel.Application"
        return self.excel


class FakeComEnvironment:
    def __init__(self) -> None:
        self.pythoncom = FakePythonCom()
        self.win32 = FakeWin32()

    @property
    def excel(self) -> FakeExcel:
        return self.win32.excel

    @property
    def workbook(self) -> FakeWorkbook:
        assert self.excel.Workbooks.last_workbook is not None
        return self.excel.Workbooks.last_workbook

    @property
    def query_sheet(self) -> FakeWorksheet:
        return self.workbook.Worksheets.query_sheet

    @property
    def result_sheet(self) -> FakeWorksheet:
        return self.workbook.Worksheets.result_sheet

    @property
    def schema_sheet(self) -> FakeWorksheet:
        return self.workbook.Worksheets.schema_sheet


def _install_fake_com(monkeypatch) -> FakeComEnvironment:
    env = FakeComEnvironment()
    monkeypatch.setattr(excel_com, "_require_excel_modules", lambda: (env.pythoncom, env.win32))
    return env


def _table_names(path: Path) -> set[str]:
    workbook = load_workbook(path)
    names: set[str] = set()
    for worksheet in workbook.worksheets:
        names.update(worksheet.tables.keys())
    return names


def test_com_create_reuses_openpyxl_ui_and_installs_power_query(tmp_path: Path, monkeypatch) -> None:
    env = _install_fake_com(monkeypatch)
    output = tmp_path / "client.xlsx"

    result = excel_com.create_workbook(output, overwrite=False, visible=True)

    assert result == output.resolve()
    assert env.excel.Visible is True
    assert env.excel.Workbooks.opened_paths == [str(output.resolve())]
    assert env.pythoncom.initialized == 1
    assert env.pythoncom.uninitialized == 1
    assert env.workbook.saved is True
    assert env.excel.quit_called is True

    workbook = load_workbook(output)
    assert workbook.sheetnames == list(client_sheet_names().all)
    assert workbook[client_sheet_names().schema].sheet_state == "hidden"
    assert set(REQUIRED_TABLE_NAMES).issubset(_table_names(output))
    assert workbook["Trino Result"]["A5"].value == (
        "В preview-книге на Unix есть только UI-таблицы. "
        "Для встраивания Power Query используйте WindowsOS."
    )

    assert [query["Name"] for query in env.workbook.Queries.added] == list(M_QUERIES)
    assert env.schema_sheet.ranges["A1:Z500"].cleared is True
    assert env.schema_sheet.ListObjects.add_calls[0]["destination"].address == "A1"
    schema_table = env.schema_sheet.ListObjects.result
    assert schema_table.Name == "tblTrinoResultSchema"
    assert schema_table.QueryTable.CommandText == "SELECT * FROM [TrinoResultSchema]"
    assert schema_table.QueryTable.RefreshStyle == excel_com.XL_OVERWRITE_CELLS
    assert env.schema_sheet.Visible == excel_com.XL_SHEET_HIDDEN
    assert env.workbook.Names.added == [
        {"Name": "TrinoSchemaColumnName", "RefersTo": "='Trino Schema'!$A:$A"},
        {"Name": "TrinoSchemaExcelFormat", "RefersTo": "='Trino Schema'!$C:$C"},
    ]
    assert env.result_sheet.ranges["A5:Z500"].cleared is True
    assert env.result_sheet.ListObjects.add_calls[0]["destination"].address == "A5"
    result_table = env.result_sheet.ListObjects.result
    assert result_table.Name == "tblTrinoResult"
    assert result_table.QueryTable.CommandType == excel_com.XL_CMD_SQL
    assert result_table.QueryTable.CommandText == "SELECT * FROM [TrinoResult]"
    assert result_table.QueryTable.BackgroundQuery is False
    assert result_table.QueryTable.RefreshOnFileOpen is False
    assert result_table.QueryTable.RefreshPeriod == 0
    assert result_table.QueryTable.EnableRefresh is True
    assert result_table.QueryTable.SaveData is True
    assert result_table.TableStyle == "TableStyleMedium4"
    assert result_table.QueryTable.PreserveFormatting is True
    assert result_table.QueryTable.PreserveColumnInfo is True
    assert result_table.QueryTable.AdjustColumnWidth is True
    assert result_table.QueryTable.RefreshStyle == excel_com.XL_OVERWRITE_CELLS
    format_conditions = env.result_sheet.ranges[excel_com.RESULT_TEMPORAL_FORMAT_RANGE].FormatConditions.items
    assert [condition.NumberFormat for condition in format_conditions] == [
        "dd.mm.yyyy",
        "dd.mm.yyyy hh:mm",
        "hh:mm:ss",
    ]
    assert "TrinoSchemaExcelFormat" in format_conditions[0].Formula1
    assert "LOCAL_INDEX" in format_conditions[0].Formula1
    assert "A$5" in format_conditions[0].Formula1
    assert env.result_sheet.ranges["A6"].cleared is True
    assert env.workbook.schema_connection.RefreshWithRefreshAll is True
    assert env.workbook.schema_connection.RefreshOnFileOpen is False
    assert env.workbook.result_connection.RefreshWithRefreshAll is True
    assert env.workbook.result_connection.RefreshOnFileOpen is False
    assert env.workbook.result_connection.OLEDBConnection.BackgroundQuery is False
    assert env.workbook.result_connection.OLEDBConnection.RefreshPeriod == 0
    assert env.workbook.helper_connection.RefreshWithRefreshAll is False
    assert env.workbook.helper_connection.RefreshOnFileOpen is False


def test_com_install_preserves_existing_workbook_content(tmp_path: Path, monkeypatch) -> None:
    env = _install_fake_com(monkeypatch)
    source = tmp_path / "source.xlsx"
    output = tmp_path / "installed.xlsx"

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Existing"
    sheet["A1"] = "keep me"
    workbook.save(source)

    result = excel_com.install_client(source, output_path=output, overwrite=False, visible=False)

    assert result == output.resolve()
    assert env.excel.Workbooks.opened_paths == [str(output.resolve())]
    installed = load_workbook(output)
    assert installed["Existing"]["A1"].value == "keep me"
    assert set(client_sheet_names().all).issubset(installed.sheetnames)
    assert [query["Name"] for query in env.workbook.Queries.added] == list(M_QUERIES)


def test_com_create_keeps_workbook_when_result_table_creation_fails(tmp_path: Path, monkeypatch) -> None:
    env = _install_fake_com(monkeypatch)
    output = tmp_path / "client.xlsx"

    original_open = env.excel.Workbooks.Open

    def open_with_result_table_failure(path: str) -> FakeWorkbook:
        workbook = original_open(path)
        workbook.Worksheets.result_sheet.ListObjects.fail_add = True
        return workbook

    env.excel.Workbooks.Open = open_with_result_table_failure  # type: ignore[method-assign]

    result = excel_com.create_workbook(output, overwrite=False, visible=False)

    assert result == output.resolve()
    assert output.exists()
    assert env.result_sheet.ranges["A5"].Value.startswith("Power Query-запросы добавлены")
    assert [query["Name"] for query in env.workbook.Queries.added] == list(M_QUERIES)


def test_com_create_advanced_xlsm_installs_vba_and_button(tmp_path: Path, monkeypatch) -> None:
    env = _install_fake_com(monkeypatch)
    output = tmp_path / "advanced.xlsm"

    result = excel_com.create_workbook(output, overwrite=False, visible=False, advanced=True)

    assert result == output.resolve()
    assert env.workbook.save_as_calls[0]["FileFormat"] == excel_com.XL_OPENXML_WORKBOOK_MACRO_ENABLED
    assert env.workbook.VBProject.VBComponents.added
    component = env.workbook.VBProject.VBComponents.added[0]
    assert component.Name == "TrinoAdvancedClient"
    assert "TrinoResultSchema" not in [query["Name"] for query in env.workbook.Queries.added]
    assert "TrinoRunQueryToSheet" in component.CodeModule.code
    assert "tblTrinoAdvancedTarget" in component.CodeModule.code
    assert "BuildFriendlyErrorMessage" in component.CodeModule.code
    assert "ResourceAccessForbidden" in component.CodeModule.code
    assert "PrepareResultWorksheet resultWs\n    UpsertResultQuery queryName, sqlText" in component.CodeModule.code
    assert 'lo.TableStyle = "TableStyleMedium4"' in component.CodeModule.code
    assert ".PreserveFormatting = True" in component.CodeModule.code
    assert ".PreserveColumnInfo = True" in component.CodeModule.code
    assert ".RefreshStyle = 0" in component.CodeModule.code
    assert "PrepareResultWorksheet resultWs" in component.CodeModule.code
    assert "IsManagedResultTableName" in component.CodeModule.code
    assert "Private Function IsManagedResultTableName" in component.CodeModule.code
    assert "IsManagedResultTableName = (Left$(tableName, Len(\"tblTrinoResult\")) = \"tblTrinoResult\")\nEnd Function" in component.CodeModule.code
    assert "BuildResultFormula = _" not in component.CodeModule.code

    button = env.query_sheet.Shapes.items[0]
    assert button.Name == "btnTrinoRunQueryToSheet"
    assert button.OnAction == "'advanced.xlsm'!TrinoAdvancedClient.TrinoRunQueryToSheet"
    assert button.TextFrame.Characters().Text == "Выгрузить в лист"
    assert env.result_sheet.ListObjects.add_calls == []

    workbook = load_workbook(output, keep_vba=True)
    assert "tblTrinoAdvancedTarget" in _table_names(output)
    assert workbook["Trino Query"]["A12"].value == "target_sheet"
